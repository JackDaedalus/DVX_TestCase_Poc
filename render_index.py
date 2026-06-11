# -*- coding: utf-8 -*-
"""Render the suite-level overview / epic document (Word + Markdown).

This is the entry-point document for the test suite: it summarises the epic,
scope, environment, and provides a traceability matrix linking every User Story
to its functional area and test-case coverage.
"""

from __future__ import annotations

import datetime as _dt

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

import theme as T
from render_docx import (
    _rgb, _style_base, _set_cell_bg, _set_cell_margins, _set_table_borders,
    _no_table_borders, _disable_autofit_fixed, _set_col_widths, _section_heading,
    _bulleted_callout, _callout, _add_top_rule, _set_repeat_header,
    _set_cell_vertical_alignment, _add_bottom_rule, _chip_run, _left_accent_border,
)


# ---------------------------------------------------------------------------
# Word overview
# ---------------------------------------------------------------------------
def build_index_docx(suite, stories, out_path):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    _style_base(doc)

    # Footer
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_top_rule(fp, T.BORDER, size=6)
    fr = fp.add_run("DXV Feedback (Quarantine Export) \u2014 Test Suite Overview   \u2022   {}".format(
        suite["epic_key"]))
    fr.font.size = Pt(8)
    fr.font.color.rgb = _rgb(T.MUTED)

    # Title band (taller, epic styling)
    band = doc.add_table(rows=1, cols=1)
    _no_table_borders(band)
    _disable_autofit_fixed(band)
    cell = band.rows[0].cells[0]
    _set_cell_bg(cell, T.PRIMARY)
    _set_cell_margins(cell, top=240, bottom=240, left=240, right=240)
    pk = cell.paragraphs[0]
    pk.paragraph_format.space_after = Pt(2)
    rk = pk.add_run("EPIC TEST SUITE  \u2022  {}".format(suite["epic_key"]))
    rk.font.size = Pt(10.5); rk.font.bold = True; rk.font.color.rgb = _rgb("AEB9DA")
    pt = cell.add_paragraph(); pt.paragraph_format.space_after = Pt(2)
    rt = pt.add_run(suite["product"]); rt.font.size = Pt(22); rt.font.bold = True
    rt.font.color.rgb = _rgb("FFFFFF")
    ps = cell.add_paragraph()
    rs = ps.add_run(suite["module"]); rs.font.size = Pt(12); rs.font.color.rgb = _rgb("C9D2EC")
    rs.italic = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # Quick facts grid
    facts = [
        ("Pipeline", "`{}`".replace("`", "").format(suite["pipeline"]), "Epic Key", suite["epic_key"]),
        ("Version", suite["version"], "Author", suite["author"]),
        ("User Stories", str(len(stories)), "Total Test Cases",
         str(sum(len(s["tests"]) for s in stories))),
        ("Total Test Steps", str(sum(len(t["steps"]) for s in stories for t in s["tests"])),
         "Generated", _dt.date.today().isoformat()),
    ]
    tbl = doc.add_table(rows=len(facts), cols=4)
    _disable_autofit_fixed(tbl)
    _set_table_borders(tbl, color="FFFFFF", size=8)
    _set_col_widths(tbl, [1.6, 2.0, 1.7, 1.85])
    for r, (k1, v1, k2, v2) in enumerate(facts):
        cs = tbl.rows[r].cells
        data = [(k1, True), (v1, False), (k2, True), (v2, False)]
        for ci, (text, is_label) in enumerate(data):
            c = cs[ci]
            _set_cell_margins(c, top=60, bottom=60, left=120, right=120)
            _set_cell_vertical_alignment(c, "center")
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            if is_label:
                _set_cell_bg(c, T.PRIMARY_LIGHT)
                run = p.add_run(text); run.font.bold = True; run.font.size = Pt(9)
                run.font.color.rgb = _rgb("FFFFFF")
            else:
                _set_cell_bg(c, T.ZEBRA_FILL if r % 2 == 0 else "F6F8FC")
                run = p.add_run(text); run.font.size = Pt(9.5)
                run.font.color.rgb = _rgb("23272F")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # Summary
    _section_heading(doc, "Suite Summary")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(suite["summary"]); r.font.size = Pt(10.5)

    # Scope
    _section_heading(doc, "In Scope")
    _bulleted_callout(doc, suite["in_scope"], T.AC_FILL, T.AC_BAR, "Covered by this suite")
    _section_heading(doc, "Out of Scope")
    _bulleted_callout(doc, suite["out_of_scope"], T.WARN_FILL, T.WARN_BAR, "Explicitly excluded")

    # Environment
    _section_heading(doc, "Test Environment")
    et = doc.add_table(rows=len(suite["env"]), cols=2)
    _disable_autofit_fixed(et)
    _set_table_borders(et, color=T.BORDER, size=4)
    _set_col_widths(et, [2.3, 4.85])
    for i, (k, v) in enumerate(suite["env"]):
        c0, c1 = et.rows[i].cells
        _set_cell_bg(c0, "E8ECF6" if i % 2 == 0 else "F2F5FB")
        _set_cell_bg(c1, "FFFFFF" if i % 2 == 0 else "FAFBFE")
        for c in (c0, c1):
            _set_cell_margins(c, top=50, bottom=50, left=120, right=120)
        r0 = c0.paragraphs[0].add_run(k); r0.font.bold = True; r0.font.size = Pt(9.5)
        r0.font.color.rgb = _rgb(T.SECONDARY)
        r1 = c1.paragraphs[0].add_run(v); r1.font.size = Pt(9.5)

    # Traceability matrix
    _section_heading(doc, "User Story Index & Traceability", color=T.PRIMARY)
    mt = doc.add_table(rows=1, cols=5)
    _disable_autofit_fixed(mt)
    _set_table_borders(mt, color=T.BORDER, size=4)
    _set_col_widths(mt, [1.0, 2.55, 1.85, 0.45, 0.85])
    headers = ["Story Key", "Functional Area / Title", "Epic Theme", "TCs", "Priority"]
    aligns = ["LEFT", "LEFT", "LEFT", "CENTER", "CENTER"]
    hdr = mt.rows[0]; _set_repeat_header(hdr)
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        _set_cell_bg(c, T.HEADER_FILL)
        _set_cell_margins(c, top=55, bottom=55, left=70, right=70)
        _set_cell_vertical_alignment(c, "center")
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if aligns[i] == "CENTER" else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(h); run.font.bold = True; run.font.size = Pt(8.5)
        run.font.color.rgb = _rgb(T.HEADER_TEXT)
    for idx, s in enumerate(stories):
        row = mt.add_row()
        zebra = T.ZEBRA_FILL if idx % 2 == 0 else "FFFFFF"
        cells = row.cells
        for c in cells:
            _set_cell_bg(c, zebra)
            _set_cell_margins(c, top=50, bottom=50, left=70, right=70)
            _set_cell_vertical_alignment(c, "center")
        rk = cells[0].paragraphs[0].add_run(s["id"]); rk.font.bold = True; rk.font.size = Pt(8.5)
        rk.font.color.rgb = _rgb(T.PRIMARY_LIGHT)
        rt = cells[1].paragraphs[0].add_run(s["title"]); rt.font.size = Pt(8.5)
        re = cells[2].paragraphs[0].add_run(s["epic"]); re.font.size = Pt(8)
        re.font.color.rgb = _rgb(T.SECONDARY)
        pc = cells[3].paragraphs[0]; pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rc = pc.add_run(str(len(s["tests"]))); rc.font.size = Pt(9); rc.font.bold = True
        rc.font.color.rgb = _rgb(T.PRIMARY_LIGHT)
        pp = cells[4].paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if s["priority"] in T.PRIORITY:
            pf, pcol = T.PRIORITY[s["priority"]]
            _chip_run(pp, s["priority"], pf, pcol)
        else:
            rp = pp.add_run(s["priority"]); rp.font.size = Pt(8)
            rp.font.color.rgb = _rgb(T.MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # How to use
    _section_heading(doc, "How to Use This Suite")
    _callout(doc, [
        "Each User Story is delivered as a standalone Word (.docx) document and a parallel "
        "Markdown (.md) file with identical content. The Markdown files carry YAML front "
        "matter and are intended for automated import into Atlassian JIRA Cloud.",
        "Execute the test cases in the order presented within each story. 'Happy Path' cases "
        "establish baseline behaviour; 'Negative', 'Boundary', 'Security' and 'Non-Functional' "
        "cases probe edges, failure handling and quality attributes.",
        "Substitute the example test-data values (process_date, table names, paths, thresholds) "
        "with values appropriate to your environment and SLAs before execution.",
    ], T.INFO_FILL, T.INFO_BAR)

    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10)
    _add_top_rule(p, T.BORDER, size=6)
    note = p.add_run("Source: feedback-quarantine-export.md  \u2014  generated {}.".format(
        _dt.date.today().isoformat()))
    note.font.size = Pt(8); note.italic = True; note.font.color.rgb = _rgb(T.MUTED)

    doc.save(out_path)
    return out_path


# ---------------------------------------------------------------------------
# Markdown overview
# ---------------------------------------------------------------------------
def build_index_md(suite, stories, out_path):
    L = []
    L.append("---")
    L.append('jira_issue_type: "Epic"')
    L.append('epic_key: "{}"'.format(suite["epic_key"]))
    L.append('epic_name: "{} \u2014 {}"'.format(suite["product"], suite["module"]))
    L.append('version: "{}"'.format(suite["version"]))
    L.append("user_story_count: {}".format(len(stories)))
    L.append("total_test_cases: {}".format(sum(len(s["tests"]) for s in stories)))
    L.append("---")
    L.append("")
    L.append("# {} \u2014 Test Suite Overview".format(suite["product"]))
    L.append("")
    L.append("**Module:** {}  ".format(suite["module"]))
    L.append("**Pipeline:** `{}`  |  **Epic Key:** `{}`  |  **Version:** {}".format(
        suite["pipeline"], suite["epic_key"], suite["version"]))
    L.append("")

    L.append("| Metric | Value |")
    L.append("| --- | --- |")
    L.append("| User Stories | {} |".format(len(stories)))
    L.append("| Total Test Cases | {} |".format(sum(len(s["tests"]) for s in stories)))
    L.append("| Total Test Steps | {} |".format(
        sum(len(t["steps"]) for s in stories for t in s["tests"])))
    L.append("| Author | {} |".format(suite["author"]))
    L.append("| Generated | {} |".format(_dt.date.today().isoformat()))
    L.append("")

    L.append("## Suite Summary")
    L.append("")
    L.append(suite["summary"])
    L.append("")

    L.append("## In Scope")
    L.append("")
    for it in suite["in_scope"]:
        L.append("- {}".format(it))
    L.append("")
    L.append("## Out of Scope")
    L.append("")
    for it in suite["out_of_scope"]:
        L.append("- {}".format(it))
    L.append("")

    L.append("## Test Environment")
    L.append("")
    L.append("| Aspect | Setting |")
    L.append("| --- | --- |")
    for k, v in suite["env"]:
        L.append("| **{}** | `{}` |".format(k, v))
    L.append("")

    L.append("## User Story Index & Traceability")
    L.append("")
    L.append("| Story Key | Functional Area / Title | Epic Theme | Test Cases | Priority | Document |")
    L.append("| --- | --- | --- | :---: | :---: | --- |")
    for s in stories:
        L.append("| `{}` | {} | {} | {} | {} | [`{}.md`]({}.md) |".format(
            s["id"], s["title"], s["epic"], len(s["tests"]), s["priority"],
            s["filename"], s["filename"]))
    L.append("")

    L.append("## How to Use This Suite")
    L.append("")
    L.append("- Each User Story is delivered as a standalone Word (`.docx`) document and a "
             "parallel Markdown (`.md`) file with identical content. The Markdown files carry "
             "YAML front matter and are intended for automated import into Atlassian JIRA Cloud.")
    L.append("- Execute the test cases in the order presented within each story. *Happy Path* "
             "cases establish baseline behaviour; *Negative*, *Boundary*, *Security* and "
             "*Non-Functional* cases probe edges, failure handling and quality attributes.")
    L.append("- Substitute the example test-data values (`process_date`, table names, paths, "
             "thresholds) with values appropriate to your environment and SLAs before execution.")
    L.append("")
    L.append("---")
    L.append("")
    L.append("_Source: `feedback-quarantine-export.md` \u2014 generated {}._".format(
        _dt.date.today().isoformat()))
    L.append("")

    with open(out_path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L))
    return out_path

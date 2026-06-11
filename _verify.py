import re
import glob
from docx import Document

doc = Document('DXV-FRE-001_Pipeline-Invocation-and-Parameter-Resolution.docx')
xml = doc.element.xml
fills = sorted(set(re.findall(r'w:fill="([0-9A-Fa-f]{6})"', xml)))
print('Distinct fill colours used:', fills)
print('Has dark header fill (1F2A56):', '1F2A56' in fills)
print('Has zebra fill (EEF1F8):', 'EEF1F8' in fills)
print('Has AC green tint (ECF7EE):', 'ECF7EE' in fills)
print('Has info blue tint (EAF1FB):', 'EAF1FB' in fills)
print('Number of tables:', len(doc.tables))

hdr_texts = []
for t in doc.tables:
    first_row = [c.text.strip() for c in t.rows[0].cells]
    if 'Expected Result' in first_row:
        hdr_texts.append(first_row)
print('Step tables with proper headers:', len(hdr_texts))
print('Sample header row:', hdr_texts[0] if hdr_texts else None)

sec = doc.sections[0]
ftxt = sec.footer.paragraphs[0].text
print('Footer text sample:', repr(ftxt[:70]))
htxt = sec.header.paragraphs[0].text
print('Header text sample:', repr(htxt[:70]))

# Confirm left-accent bars (callout boxes) present
has_left_bar = 'w:left' in xml and 'EAF1FB' in fills
print('Callout left-accent + tint present:', has_left_bar)

# Page numbering fields — check the FOOTER part, not the document body
footer_xml = sec.footer._element.xml
print('Footer has PAGE field:', 'PAGE' in footer_xml)
print('Footer has NUMPAGES field:', 'NUMPAGES' in footer_xml)
print('Footer has fldChar:', 'fldChar' in footer_xml)

# Cross-check all docs have a title band cell with primary fill
print('\nPer-doc title-band check:')
for d in sorted(glob.glob('DXV-FRE-*.docx')):
    x = Document(d).element.xml
    ok = '1F2A56' in x
    print('  {}  primary-band={}'.format(d[:55].ljust(55), ok))
